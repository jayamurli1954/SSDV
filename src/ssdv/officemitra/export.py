from __future__ import annotations

from pathlib import Path
from typing import Any


def _sheet_title(ws_title: str) -> str:
    # Excel sheet titles have a 31-char limit.
    return ws_title[:31]


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def write_excel(payload: dict[str, Any], out: Path) -> None:
    """Write a lightweight executive Excel workbook for dashboard payload."""

    from openpyxl import Workbook

    wb = Workbook()

    # Summary
    ws = wb.active
    ws.title = _sheet_title("Summary")
    rows = [
        ("Company", _stringify(payload.get("company"))),
        ("As-of", _stringify(payload.get("as_of"))),
        ("Scenario", _stringify(payload.get("scenario"))),
        ("Pack", _stringify(payload.get("pack"))),
        ("Title", _stringify(payload.get("title"))),
    ]
    for r_i, (k, v) in enumerate(rows, start=1):
        ws.cell(row=r_i, column=1, value=k)
        ws.cell(row=r_i, column=2, value=v)

    kpis = payload.get("kpis") or {}

    ws2 = wb.create_sheet(_sheet_title("Key KPIs"))
    kpi_rows = [
        ("Sales", _stringify(kpis.get("sales"))),
        ("COGS", _stringify(kpis.get("cogs"))),
        ("Gross Margin", _stringify(kpis.get("gross_margin"))),
        ("GM %", _stringify(kpis.get("gm_pct"))),
        ("Opex", _stringify(kpis.get("opex"))),
        ("Purchases", _stringify(kpis.get("purchases"))),
        ("Receipts", _stringify(kpis.get("receipts"))),
        ("Payments", _stringify(kpis.get("payments"))),
        ("Working capital", _stringify(kpis.get("working_capital"))),
        ("Collection efficiency", _stringify(kpis.get("collection_efficiency"))),
        ("YTD profit", _stringify(kpis.get("ytd_profit"))),
        ("Monthly profit", _stringify(kpis.get("monthly_profit"))),
    ]
    ws2.append(["Metric", "Value"])
    for k, v in kpi_rows:
        ws2.append([k, v])

    # Aging
    ws3 = wb.create_sheet(_sheet_title("Aging"))
    ws3.append(["A/R Aging bucket", "Value"])
    ar_aging = (kpis.get("ar_aging") or {}) if isinstance(kpis.get("ar_aging"), dict) else {}
    for bucket, value in ar_aging.items():
        ws3.append([bucket, _stringify(value)])

    ws3.append([])
    ws3.append(["A/P Aging bucket", "Value"])
    ap_aging = (kpis.get("ap_aging") or {}) if isinstance(kpis.get("ap_aging"), dict) else {}
    for bucket, value in ap_aging.items():
        ws3.append([bucket, _stringify(value)])

    # Parties
    ws4 = wb.create_sheet(_sheet_title("Top Parties"))
    parties = payload.get("parties") or {}
    ws4.append(["Customer", "AR 90+", "Outstanding", "Unaged", "Oldest days"])
    for row in parties.get("overdue_customers") or []:
        ws4.append(
            [
                _stringify(row.get("name")),
                _stringify(row.get("overdue_90")),
                _stringify(row.get("outstanding")),
                _stringify(row.get("unaged")),
                _stringify(row.get("oldest_days")),
            ]
        )

    ws4.append([])
    ws4.append(["Vendor", "AP 90+", "Outstanding", "Unaged", "Oldest days"])
    for row in parties.get("vendor_exposure") or []:
        ws4.append(
            [
                _stringify(row.get("name")),
                _stringify(row.get("overdue_90")),
                _stringify(row.get("outstanding")),
                _stringify(row.get("unaged")),
                _stringify(row.get("oldest_days")),
            ]
        )

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def write_ppt(payload: dict[str, Any], out: Path) -> None:
    """Write a simple executive PPTX from the dashboard payload."""

    from pptx import Presentation

    prs = Presentation()

    def _add_title_slide(title: str, lines: list[str]) -> None:
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
        slide.shapes.title.text = title
        tf = slide.shapes.placeholders[1].text_frame
        tf.clear()
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.level = 0

    kpis = payload.get("kpis") or {}
    parties = payload.get("parties") or {}

    company = _stringify(payload.get("company"))
    as_of = _stringify(payload.get("as_of"))
    scenario = _stringify(payload.get("scenario"))

    _add_title_slide(
        "OfficeMitra Board Pack",
        [
            f"Company: {company}",
            f"As-of: {as_of}",
            f"Scenario: {scenario}",
            f"Sales: {_stringify(kpis.get('sales'))}",
            f"COGS: {_stringify(kpis.get('cogs'))}",
            f"Gross Margin: {_stringify(kpis.get('gross_margin'))} ({_stringify(kpis.get('gm_pct'))})",
        ],
    )

    _add_title_slide(
        "Aging & Liquidity Signals",
        [
            f"DSO/DIO/DPO: {_stringify(kpis.get('dso'))}/{_stringify(kpis.get('dio'))}/{_stringify(kpis.get('dpo'))}",
            f"CCC: {_stringify(kpis.get('ccc'))}",
            f"Working capital: {_stringify(kpis.get('working_capital'))}",
            f"Collection efficiency: {_stringify(kpis.get('collection_efficiency'))}",
            f"Cash (monthly net): {_stringify(kpis.get('monthly_net_cash'))}",
        ],
    )

    cust_rows = parties.get("overdue_customers") or []
    vend_rows = parties.get("vendor_exposure") or []
    cust_lines = [
        f"{_stringify(r.get('name'))}: AR 90+ {_stringify(r.get('overdue_90'))}, Outstanding {_stringify(r.get('outstanding'))}"
        for r in cust_rows[:5]
    ]
    vend_lines = [
        f"{_stringify(r.get('name'))}: AP 90+ {_stringify(r.get('overdue_90'))}, Outstanding {_stringify(r.get('outstanding'))}"
        for r in vend_rows[:5]
    ]

    _add_title_slide(
        "Top overdue customers & vendor exposure",
        [
            "Customers:",
            *(cust_lines if cust_lines else ["(none)"]),
            "",
            "Vendors:",
            *(vend_lines if vend_lines else ["(none)"]),
        ],
    )

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
